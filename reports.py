"""PDF (reportlab) and Excel (openpyxl) report builders for the three tools. All return bytes."""
from __future__ import annotations

import datetime as dt
import io
import os

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import charts
from mis_engine import AMT, BUCKETS
from utils import fdate, inr, inr_short

_FD = os.path.join(mpl.get_data_path(), "fonts", "ttf")
pdfmetrics.registerFont(TTFont("DV", os.path.join(_FD, "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("DV-B", os.path.join(_FD, "DejaVuSans-Bold.ttf")))
pdfmetrics.registerFontFamily("DV", normal="DV", bold="DV-B", italic="DV", boldItalic="DV-B")

NAVY, TEALC, GREYC, LIGHT = "#1f3a5f", "#2a9d8f", "#6b7280", "#f3f4f6"
GOOD, BAD = "#2e7d32", "#c62828"


# ------------------------------------------------------------------ building blocks
def _esc(x) -> str:
    return str(x).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _style(name, size=9, bold=False, color="#111827", lead=None, align=None, space=0):
    kw = dict(fontName="DV-B" if bold else "DV", fontSize=size, leading=lead or size + 3, textColor=colors.HexColor(color), spaceAfter=space)
    if align:
        kw["alignment"] = align
    return ParagraphStyle(name, **kw)


def _p(text, size=9, bold=False, color="#111827", space=0, raw=False):
    return Paragraph(text if raw else _esc(text), _style("p", size, bold, color, space=space))


def _h(text):
    return Paragraph(_esc(text), _style("h", 11, True, NAVY, space=3))


def _t(rows, widths, right=(), size=7.2, bold_rows=(), zebra=True, color_col=None):
    sl, sr = _style("l", size, lead=size + 2.6), _style("r", size, lead=size + 2.6, align=TA_RIGHT)
    sb, sbr = _style("b", size, True, lead=size + 2.6), _style("br", size, True, lead=size + 2.6, align=TA_RIGHT)
    shl, shr = _style("hl", size, True, "#ffffff", size + 2.6), _style("hr", size, True, "#ffffff", size + 2.6, TA_RIGHT)
    data = []
    for ri, r in enumerate(rows):
        data.append([Paragraph(_esc(c), (shr if ci in right else shl) if ri == 0 else
                               (sbr if ci in right else sb) if ri in bold_rows else (sr if ci in right else sl))
                     for ci, c in enumerate(r)])
    t = Table(data, colWidths=widths, repeatRows=1)
    st = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")), ("TOPPADDING", (0, 0), (-1, -1), 2),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]
    if zebra:
        st += [("BACKGROUND", (0, i), (-1, i), colors.HexColor(LIGHT)) for i in range(2, len(rows), 2)]
    st += [("BACKGROUND", (0, i), (-1, i), colors.HexColor("#e5edf7")) for i in bold_rows]
    t.setStyle(TableStyle(st))
    return t


def _tiles(items, width):
    cells = []
    for label, value, sub, good in items:
        subc = GREYC if good is None else (GOOD if good else BAD)
        cells.append(Paragraph(f'<font size="7" color="{GREYC}">{_esc(label)}</font><br/>'
                               f'<font name="DV-B" size="12" color="{NAVY}">{_esc(value)}</font><br/>'
                               f'<font size="6.8" color="{subc}">{_esc(sub)}</font>', _style("tile", 8, lead=14)))
    t = Table([cells], colWidths=[width / len(items)] * len(items))
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(LIGHT)), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                           ("LINEAFTER", (0, 0), (-2, -1), 2, colors.white), ("TOPPADDING", (0, 0), (-1, -1), 6),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (-1, -1), 7)]))
    return t


def _img(fig, width_mm):
    buf = io.BytesIO()
    w, h = fig.get_size_inches()
    fig.savefig(buf, format="png", dpi=190, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return Image(buf, width=width_mm * mm, height=width_mm * mm * h / w)


def _header(meta, title):
    return [_p(meta.get("company", ""), 15, True, NAVY, space=1), _p(title, 11, True, "#111827", space=2),
            _p("  |  ".join(x for x in [meta.get("account"), meta.get("period"),
                                         "Generated " + dt.datetime.now().strftime("%d-%m-%Y %H:%M")] if x), 8, False, GREYC, space=6)]


def _signoff(width):
    t = Table([["Prepared by:", "", "Reviewed by:", "", "Date:", ""]], colWidths=[width * x for x in (.12, .22, .12, .22, .08, .24)])
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "DV", 8), ("LINEBELOW", (1, 0), (1, 0), .5, colors.black),
                           ("LINEBELOW", (3, 0), (3, 0), .5, colors.black), ("LINEBELOW", (5, 0), (5, 0), .5, colors.black),
                           ("TOPPADDING", (0, 0), (-1, -1), 14)]))
    return t


def _build(story, title, landscape_mode=False) -> bytes:
    buf = io.BytesIO()
    ps = landscape(A4) if landscape_mode else A4
    doc = SimpleDocTemplate(buf, pagesize=ps, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=17 * mm, bottomMargin=15 * mm,
                            title=title, author="CA Analytics Suite")

    def deco(c, d):
        c.saveState()
        c.setStrokeColor(colors.HexColor(NAVY))
        c.setLineWidth(1.2)
        c.line(14 * mm, ps[1] - 12 * mm, ps[0] - 14 * mm, ps[1] - 12 * mm)
        c.setFont("DV-B", 8)
        c.setFillColor(colors.HexColor(NAVY))
        c.drawString(14 * mm, ps[1] - 10.5 * mm, title)
        c.setFont("DV", 7)
        c.setFillColor(colors.HexColor(GREYC))
        c.drawString(14 * mm, 8 * mm, "Computer-generated working paper - review and verify before reliance.")
        c.drawRightString(ps[0] - 14 * mm, 8 * mm, f"Page {d.page}")
        c.restoreState()

    doc.build(story, onFirstPage=deco, onLaterPages=deco)
    return buf.getvalue()


def _rows(df, spec, limit=None):
    """spec: list of (header, column, formatter)."""
    out = [[h for h, _, _ in spec]]
    d = df if limit is None else df.head(limit)
    for _, r in d.iterrows():
        out.append([f(r[c]) if f else r[c] for _, c, f in spec])
    return out


def _w(total, fr):
    s = sum(fr)
    return [total * x / s for x in fr]


# ------------------------------------------------------------------ BANK RECON
def bank_recon_pdf(res, meta) -> bytes:
    W = 182 * mm
    s = res.summary
    st = _header(meta, "Bank Reconciliation Statement")
    st.append(_tiles([("Books entries", str(s["n_books"]), f"{s['books_matched']} matched", None),
                      ("Bank entries", str(s["n_bank"]), f"{s['bank_matched']} matched", None),
                      ("Entries matched", f"{s['pct_books_matched']:.1f}%", "of books entries", s["pct_books_matched"] >= 90),
                      ("Value matched", f"{s['pct_value_matched']:.1f}%", "of books value", s["pct_value_matched"] >= 90),
                      ("Open exceptions", str(s["n_exceptions"]), "need review", s["n_exceptions"] == 0)], W))
    st.append(Spacer(1, 6))
    ok = s["tie"]
    st.append(_p(f'<b>Status:</b> <font color="{GOOD if ok else BAD}"><b>{"RECONCILED" if ok else "NOT RECONCILED"}</b></font> - adjusted books '
                 f'balance {inr(s["adj_books"])} vs adjusted bank balance {inr(s["adj_bank"])} (difference {inr(s["difference"])}).',
                 8.5, raw=True, space=6))
    for sec in ["A. Adjusting the books balance", "B. Adjusting the bank balance"]:
        b = res.brs[res.brs["section"] == sec]
        rows = [[sec, "Amount (₹)"]] + [[r.particulars, inr(r.amount)] for r in b.itertuples()]
        bold = [i + 1 for i, r in enumerate(b.itertuples()) if r.strong]
        st += [_t(rows, _w(W, [78, 22]), right=(1,), size=7.8, bold_rows=bold, zebra=False), Spacer(1, 7)]
    st += [_t([["Check", "Amount (₹)"], ["Difference between adjusted balances (should be nil)", inr(s["difference"])]],
              _w(W, [78, 22]), right=(1,), size=7.8, bold_rows=[1], zebra=False), Spacer(1, 8)]
    st += [_h("Exceptions at a glance"), _img(charts.bank_exceptions_chart(res), 170), Spacer(1, 6)]

    spec = [("Date", "date", fdate), ("Ref", "ref", None), ("Description", "desc", None), ("Category", "category", None), ("Amount (₹)", "amount", inr)]
    for title, df, note in [("Entries in books but not in bank", res.books_only,
                             "Timing items (deposits in transit, uncleared cheques) reverse next period; duplicates need correction in books."),
                            ("Entries in bank but not in books", res.bank_only,
                             "Pass journal entries in books for charges, interest and direct debits.")]:
        blk = [_h(title), _p(note, 7.5, color=GREYC, space=3)]
        blk.append(_t(_rows(df, spec, 60), _w(W, [14, 12, 38, 42, 18]), right=(4,)) if len(df) else _p("None.", 8))
        st += blk + [Spacer(1, 8)]
    mm_ = res.matched[res.matched["match_type"].str.startswith("3")]
    if len(mm_):
        st += [_h("Amount differences (same reference, different amount)"),
               _t(_rows(mm_, [("Ref", "books_ref", None), ("Books date", "books_date", fdate), ("Bank date", "bank_date", fdate),
                              ("Books (₹)", "books_amount", inr), ("Bank (₹)", "bank_amount", inr), ("Difference (₹)", "diff", inr)]),
                  _w(W, [18, 16, 16, 20, 20, 20]), right=(3, 4, 5)), Spacer(1, 8)]
    st += [PageBreak(), _h("Matched transactions (appendix)"),
           _p("Pass 1 exact reference + amount; 2 amount within the date window; 3 same reference with an amount difference; "
              "4/5 grouped entries (cheques banked together or split).", 7.5, color=GREYC, space=3)]
    mt = res.matched.sort_values(["match_type", "books_date"])
    st.append(_t(_rows(mt, [("Match type", "match_type", lambda x: x.split(". ", 1)[1]), ("Books date", "books_date", fdate),
                            ("Ref", "books_ref", None), ("Bank date", "bank_date", fdate), ("Books (₹)", "books_amount", inr),
                            ("Bank (₹)", "bank_amount", inr), ("Gap (d)", "days_gap", lambda x: "" if pd.isna(x) else int(x))], 250),
                 _w(W, [34, 16, 18, 16, 20, 20, 9]), right=(4, 5, 6), size=6.8))
    if len(mt) > 250:
        st.append(_p(f"Showing first 250 of {len(mt)} matched rows; the full list is in the Excel export.", 7.5, color=GREYC))
    st += [Spacer(1, 12), _signoff(W)]
    return _build(st, "Bank Reconciliation Statement")


# ------------------------------------------------------------------ GST RECON
def gst_recon_pdf(res, meta) -> bytes:
    from recon_engine import S_FUZZY, S_GONLY, S_PRONLY, S_VALUE, S_DUP

    W = 182 * mm
    s = res.summary
    st = _header(meta, "GSTR-2B vs Purchase Register - ITC Reconciliation")
    st.append(_tiles([("ITC as per books", inr_short(s["itc_books"]), f"{s['n_pr']} invoices", None),
                      ("ITC as per GSTR-2B", inr_short(s["itc_2b"]), f"{s['n_2b']} invoices", None),
                      ("Eligible (matched)", inr_short(s["eligible"]), f"{s['pct_matched']:.0f}% invoices matched", s["pct_matched"] >= 90),
                      ("ITC at risk", inr_short(s["at_risk"]), "claimed, not supported", s["at_risk"] == 0),
                      ("In 2B, not booked", inr_short(s["g_only"]), f"{s['n_g_only']} invoices", s["g_only"] == 0)], W))
    st += [Spacer(1, 8), _h("ITC bridge"), _t([["Particulars", "Amount (₹)"],
           ["ITC as per purchase register", inr(s["itc_books"])],
           ["Less: ITC not available in GSTR-2B (invoices missing)", inr(-s["pr_only"])],
           ["Less: excess over GSTR-2B on value mismatches", inr(-s["excess"])],
           ["Less: duplicate entries in purchase register", inr(-s["dup"])],
           ["ITC supported by GSTR-2B (eligible)", inr(s["eligible"])],
           ["Add: ITC in GSTR-2B not yet booked", inr(s["g_only"])],
           ["Add: GSTR-2B higher than books on matched invoices", inr(s["g_excess"])],
           ["ITC as per GSTR-2B", inr(s["itc_2b"])]], _w(W, [78, 22]), right=(1,), size=7.8, bold_rows=[1, 5, 8], zebra=False),
           _p(f"Checks (nil expected): books side {inr(s['check_diff'])}; GSTR-2B side {inr(s['check_2b'])}.", 7, color=GREYC, space=6),
           _img(charts.gst_itc_chart(res), 170), _img(charts.gst_status_chart(res), 150), Spacer(1, 6)]

    sup = res.suppliers
    st += [_h("Supplier follow-up list"), _p("Suppliers to chase for filing/amending GSTR-1, or invoices to book.", 7.5, color=GREYC, space=3)]
    st.append(_t(_rows(sup, [("GSTIN", "gstin", None), ("Supplier", "supplier", None), ("Issues", "issues", None),
                             ("ITC at risk (₹)", "itc_at_risk", inr), ("In 2B not booked (₹)", "in_2b_not_in_books", inr),
                             ("Action", "action", None)], 25), _w(W, [24, 28, 32, 18, 20, 30]), right=(3, 4), size=6.8)
              if len(sup) else _p("No supplier issues.", 8))

    d = res.detail
    spec_inv = [("Supplier", "supplier", lambda x: str(x)[:22]), ("Invoice (books)", "pr_invoice", None), ("Invoice (2B)", "g_invoice", None),
                ("Taxable books", "pr_taxable", inr), ("Taxable 2B", "g_taxable", inr), ("Tax books", "pr_tax", inr),
                ("Tax 2B", "g_tax", inr), ("Remark", "remark", lambda x: str(x)[:90])]
    for title, mask in [("Value mismatches and probable matches", d["status"].isin([S_VALUE, S_FUZZY])),
                        ("In books but not in GSTR-2B (ITC at risk)", d["status"] == S_PRONLY),
                        ("Duplicates in purchase register", d["status"] == S_DUP),
                        ("In GSTR-2B but not in books", d["status"] == S_GONLY)]:
        sub = d[mask]
        st += [Spacer(1, 8), KeepTogether([_h(f"{title} ({len(sub)})")])]
        st.append(_t(_rows(sub, spec_inv, 50), _w(W, [26, 18, 18, 17, 17, 15, 15, 40]), right=(3, 4, 5, 6), size=6.6) if len(sub) else _p("None.", 8))
    st += [Spacer(1, 8), _h("Data-quality flags"),
           _p(f"Invalid GSTIN format: {s['n_invalid_gstin']} record(s)  |  Period mismatches: {s['n_period_mismatch']}  |  "
              f"Duplicate invoices in books: {s['n_dup']}", 8, space=4),
           _p("GSTIN check is a format check only (not a live validation). Matching: GSTIN + normalised invoice number first, then same supplier with "
              "same taxable value and similar invoice number or date. ITC availability depends on Section 16 conditions (including reflection in "
              "GSTR-2B, payment to supplier and time limits) - confirm before reversing or claiming.", 7.5, color=GREYC),
           Spacer(1, 10), _signoff(W)]
    return _build(st, "GSTR-2B vs Purchase Register")


# ------------------------------------------------------------------ MIS
def mis_pdf(r, meta) -> bytes:
    ps = landscape(A4)
    W = ps[0] - 28 * mm
    label = f"{r.months[0].strftime('%b-%y')} to {r.months[-1].strftime('%b-%y')}"
    meta = dict(meta, period=meta.get("period") or label)
    st = _header(meta, "Management Information System (MIS) Report")
    k = r.kpis.set_index("metric")

    def tile(m, kind="amt"):
        c, ch = k.loc[m, "current"], k.loc[m, "change"]
        v = inr_short(c) if kind == "amt" else f"{c:.1f}%"
        if ch != ch:
            return (m, v, label, None)
        return (m, v, f"{'▲' if ch >= 0 else '▼'} {abs(ch):.1f}{'%' if kind == 'amt' else ' pp'} vs prior", ch >= 0)

    st.append(_tiles([tile("Revenue"), tile("GP %", "pct"), tile("EBITDA"), tile("EBITDA %", "pct"), tile("PBT")], W))
    st += [Spacer(1, 7), _h("Key observations")]
    st += [Paragraph("•&nbsp;&nbsp;" + _esc(t), _style("b", 8.3, lead=11.5, space=2)) for t in r.insights]

    st += [Spacer(1, 8), _h("Monthly profit & loss (₹ lakh)")]
    pnl = r.pnl.loc[r.months[0]:r.months[-1]].iloc[-12:]
    lines = ["Revenue", "COGS", "Gross Profit", "GP %", "Opex", "EBITDA", "EBITDA %", "Other Income", "Depreciation", "Finance Cost", "PBT"]
    tot = r.pnl.loc[r.months[0]:r.months[-1], AMT].sum()
    hdr = ["₹ lakh"] + [p.strftime("%b-%y") for p in pnl.index] + ["Total"]
    rows = [hdr]
    for ln in lines:
        if ln.endswith("%"):
            base = {"GP %": "Gross Profit", "EBITDA %": "EBITDA"}[ln]
            rows.append([ln] + [f"{v:.1f}%" for v in pnl[ln]] + [f"{tot[base] / tot['Revenue'] * 100:.1f}%" if tot["Revenue"] else ""])
        else:
            rows.append([ln] + [f"{v / 1e5:,.1f}" for v in pnl[ln]] + [f"{tot[ln] / 1e5:,.1f}"])
    ncol = len(hdr) - 1
    bold = [i for i, ln in enumerate(["hdr"] + lines) if ln in ("Gross Profit", "EBITDA", "PBT")]
    st += [_t(rows, [W * .12] + [W * .88 / ncol] * ncol, right=tuple(range(1, ncol + 1)), size=6.8, bold_rows=bold), Spacer(1, 8)]

    st += [PageBreak(), _h("Charts"), Table([[_img(charts.revenue_trend_chart(r), 128), _img(charts.cost_mix_chart(r), 128)],
                                [_img(charts.customer_chart(r), 128), _img(charts.ageing_chart(r), 128)]], colWidths=[W / 2] * 2)]

    if r.variance is not None:
        st += [PageBreak(), _h("Budget vs actual"), _t(
            [["Line", "Actual (₹)", "Budget (₹)", "Variance (₹)", "Variance %", "Status"]] +
            [[x.line, inr(x.actual, 0), inr(x.budget, 0), inr(x.variance, 0), f"{x.var_pct:+.1f}%", x.status] for x in r.variance.itertuples()],
            _w(W * .7, [20, 20, 20, 20, 12, 14]), right=(1, 2, 3, 4), size=7.5), Spacer(1, 8)]
    else:
        st.append(PageBreak())
    if len(r.anomalies):
        st += [_h("Cost spikes to review"), _t(
            [["Head", "Month", "Amount (₹)", "Typical (₹)", "x typical"]] +
            [[a.head, a.month.strftime("%b-%y"), inr(a.amount, 0), inr(a.typical, 0), f"{a.times_typical:.1f}x"] for a in r.anomalies.itertuples()],
            _w(W * .6, [26, 10, 16, 16, 10]), right=(2, 3, 4), size=7.5), Spacer(1, 8)]
    if r.ageing:
        by = r.ageing["by_customer"].head(12)
        st += [_h(f"Receivables ageing by customer (₹ lakh) - as of {r.asof.strftime('%d-%m-%Y')}"), _t(
            [["Customer"] + BUCKETS + ["Total"]] + [[i] + [f"{v / 1e5:,.1f}" for v in row[BUCKETS + ['Total']]] for i, row in by.iterrows()] +
            [["Total"] + [f"{v / 1e5:,.1f}" for v in r.ageing["totals"]] + [f"{r.ageing['total'] / 1e5:,.1f}"]],
            _w(W, [30] + [14] * 6), right=tuple(range(1, 7)), size=7.3, bold_rows=[len(by) + 1]), Spacer(1, 8)]
    if r.forecast is not None:
        st += [_h("Revenue outlook (linear trend, indicative)"), _t(
            [["Month", "Forecast (₹)", "Low (₹)", "High (₹)"]] +
            [[f.month.strftime("%b-%y"), inr(f.forecast, 0), inr(f.low, 0), inr(f.high, 0)] for f in r.forecast.itertuples()],
            _w(W * .5, [16, 20, 20, 20]), right=(1, 2, 3), size=7.5)]
    st += [Spacer(1, 6), _p("EBITDA = Gross profit - operating expenses (excludes other income). PBT = EBITDA + other income - depreciation - finance cost. "
                            "Anomalies use a robust (median/MAD) test; forecast is a simple linear trend that ignores seasonality.", 7, color=GREYC)]
    return _build(st, "MIS Report", landscape_mode=True)


# ------------------------------------------------------------------ Excel
def _polish(ws):
    from openpyxl.styles import Alignment, Font, PatternFill

    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col[:200])
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 46)
        for c in col[1:]:
            if isinstance(c.value, bool):
                continue
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
            elif isinstance(c.value, int):
                c.number_format = "#,##0"
            elif isinstance(c.value, (dt.datetime, dt.date)):
                c.number_format = "dd-mm-yyyy"
    ws.freeze_panes = "A2"


def _xlsx(sheets: dict, extra=None) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="dd-mm-yyyy") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            _polish(xw.sheets[name[:31]])
        if extra:
            extra(xw)
    return buf.getvalue()


def bank_recon_excel(res, meta) -> bytes:
    s = res.summary
    summ = pd.DataFrame({"Metric": ["Entity", "Account", "Period", "Books entries", "Bank entries", "% books entries matched",
                                    "% value matched", "Open exceptions", "Reconciled?", "Difference (₹)"],
                         "Value": [meta.get("company", ""), meta.get("account", ""), meta.get("period", ""), s["n_books"], s["n_bank"],
                                   round(s["pct_books_matched"], 1), round(s["pct_value_matched"], 1), s["n_exceptions"],
                                   "Yes" if s["tie"] else "No", float(s["difference"])]})
    return _xlsx({"Summary": summ, "BRS": res.brs.drop(columns="strong").rename(columns={"particulars": "Particulars", "amount": "Amount"}),
                  "Matched": res.matched, "Books_not_in_bank": res.books_only, "Bank_not_in_books": res.bank_only})


def gst_recon_excel(res, meta) -> bytes:
    from recon_engine import S_FUZZY, S_GONLY, S_PRONLY, S_VALUE
    s = res.summary
    summ = pd.DataFrame({"Metric": ["ITC as per books", "ITC as per GSTR-2B", "Eligible (matched)", "Excess over 2B (value mismatch)",
                                    "In books, not in 2B", "Duplicates in books", "ITC at risk (total)", "In 2B, not in books"],
                         "Amount": [s["itc_books"], s["itc_2b"], s["eligible"], s["excess"], s["pr_only"], s["dup"], s["at_risk"], s["g_only"]]})
    d = res.detail
    return _xlsx({"Summary": summ, "Detail": d, "Supplier_FollowUp": res.suppliers,
                  "Not_in_2B": d[d["status"] == S_PRONLY], "Not_in_Books": d[d["status"] == S_GONLY],
                  "Mismatches": d[d["status"].isin([S_VALUE, S_FUZZY])]})


def mis_excel(r, meta) -> bytes:
    from openpyxl.chart import BarChart, Reference

    pnl = r.pnl.copy()
    pnl.insert(0, "Month", [p.strftime("%b-%y") for p in pnl.index])
    k = r.kpis.copy()
    fact = r.ledger.copy()
    fact["fy"] = (fact["date"] - pd.offsets.MonthBegin(3)).dt.year.map(lambda y: f"FY{y}-{str(y + 1)[2:]}")
    fact["month"] = fact["month"].astype(str)
    fact["signed_amount"] = np.where(fact["group"].isin(["Revenue", "Other Income"]), fact["amount"], -fact["amount"])
    dim = pd.DataFrame({"month": [str(m) for m in r.pnl.index], "month_start": [m.to_timestamp() for m in r.pnl.index]})
    sheets = {"KPIs": k, "Monthly_PnL": pnl.reset_index(drop=True)}
    if r.variance is not None:
        sheets["Budget_Variance"] = r.variance
    sheets["Opex_Trend"] = r.opex.rename(columns=str).reset_index()
    sheets["Customers"] = r.customers.reset_index().rename(columns={"index": "customer"})
    if len(r.anomalies):
        a = r.anomalies.copy()
        a["month"] = a["month"].astype(str)
        sheets["Cost_Spikes"] = a
    if r.ageing:
        sheets["Ageing"] = r.ageing["by_customer"].reset_index()
        sheets["Open_Invoices"] = r.ageing["detail"].drop(columns=["bucket"]).assign(bucket=r.ageing["detail"]["bucket"].astype(str))
    sheets["PowerBI_Fact"] = fact
    sheets["Dim_Month"] = dim

    def add_chart(xw):
        ws = xw.sheets["Monthly_PnL"]
        n = len(pnl)
        ch = BarChart()
        ch.title, ch.height, ch.width = "Revenue vs EBITDA", 8, 18
        for col in (2, 6):
            ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
        ws.add_chart(ch, "N2")

    return _xlsx(sheets, add_chart)
